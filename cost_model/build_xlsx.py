"""Build the live Excel calculator from the formula catalog + data files.

  python build_xlsx.py [out.xlsx] [--rows reference]

Every derived cell is a formula transpiled from formulas.py, so the workbook
recalculates when any blue input changes. --rows reference writes only the
reference scenarios (used by parity.py for a fast check).
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from dsl import names_in, to_xl, toposort
from formulas import INPUTS, LLM, OCR, OCR_INPUTS
from inputs import ROW_KEYS, Data, meta, val

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
GROUP_FILL = PatternFill("solid", fgColor="D9E1F2")
YELLOW = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))

LEVER_NAMES = ("s_bh", "s_itc", "s_spare", "s_years")
SEL = ("sel_model", "sel_precision", "sel_gpu", "sel_workload", "sel_users", "sel_speed")


def numfmt(unit, name=""):
    u = unit.lower()
    if name.endswith("_usd") or u.startswith("usd"):
        return '$#,##0;($#,##0);"-"'
    if u.startswith("inr"):
        return '"₹"#,##0;("₹"#,##0);"-"'
    if u.startswith("frac"):
        return '0.0%;(0.0%);"-"'
    if u in ("0/1",):
        return "0"
    if u in ("s", "s/seq"):
        return "0.0000"
    if u in ("x",):
        return "0.000"
    if u in ("gb", "gb/gpu", "tb", "kw", "kva", "tb/s"):
        return "#,##0.0"
    return "#,##0.##"


def name(wb, nm, sheet, cell):
    ref = f"'{sheet}'!${cell[0]}${cell[1]}" if isinstance(cell, tuple) else f"'{sheet}'!{cell}"
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w


def src_text(m):
    parts = [m.get("source", ""), m.get("date", ""), m.get("note", "")]
    return " | ".join(p for p in parts if p)


class Builder:
    def __init__(self, data, rows_mode="full"):
        self.d = data
        self.rows_mode = rows_mode
        self.wb = Workbook()
        self.sources = []  # (sheet, item, value, source)
        self.global_names = set(data.globals) | set(LEVER_NAMES)
        self.row_inputs = [k for k in INPUTS if k not in self.global_names]
        self.llm_order = toposort(LLM, set(INPUTS))
        self.ocr_order = toposort(OCR, set(INPUTS) | set(OCR_INPUTS))

    # ------------------------------------------------------------------ inputs
    def sheet_readme(self):
        ws = self.wb.active
        ws.title = "Read me"
        lines = [
            ("Self-hosting LLM cost model: Chandra OCR 2, GLM-5.3, DeepSeek V4 Pro, Kimi K2.7 Code", BOLD),
            ("Each model is sized as a standalone deployment for 50-90 concurrent users. Pune (Baner), India. INR unless marked USD.", BLACK),
            ("", BLACK),
            ("How to use", BOLD),
            ("1. 'Inputs' holds every assumption (blue = editable). The yellow block at the top holds the scenario levers: business hours, GST credit, spare node, years.", BLACK),
            ("2. 'Calculator' sizes and prices one configuration: pick model, precision, GPU, workload, users and speed from the yellow dropdowns.", BLACK),
            ("3. 'Matrix' runs every model x precision x GPU x workload x users x speed combination with the same formulas. 'Summary' picks the cheapest feasible option per case.", BLACK),
            ("4. 'Models', 'GPUs', 'Efficiency', 'Workloads' and 'OCR inputs' are lookup tables; edit any blue cell and everything recalculates.", BLACK),
            ("5. 'Sources' lists where every number came from, with dates. Rows marked ASSUMPTION have no published source.", BLACK),
            ("", BLACK),
            ("Colour legend", BOLD),
            ("Blue text: hard-coded input. Black: formula. Green: lookup from another sheet. Yellow fill: scenario lever / selection.", BLACK),
            ("", BLACK),
            ("Key method notes", BOLD),
            ("Memory: weights + KV per GPU under DP-attention + expert parallelism; KV sized at the 99th percentile of summed context.", BLACK),
            ("Speed: decode step = weight read + collective latency + per-sequence (KV read + compute + chunked prefill share); batch per replica is the largest that meets the target tok/s.", BLACK),
            ("Context: 1M-token ceiling with auto-summarisation. Context saw-tooths between the post-summary floor and the trigger T; Kimi K2.7 is capped by its 262,144 native window.", BLACK),
            ("FP32 rows are memory/cost only: no serving engine runs these models in FP32 and it adds no quality (marked infeasible).", BLACK),
            ("Capex is converted to a monthly charge with a capital recovery factor (cost of capital, resale value), not straight-line.", BLACK),
        ]
        for i, (t, f) in enumerate(lines, 1):
            c = ws.cell(row=i, column=1, value=t)
            c.font = f
        ws.column_dimensions["A"].width = 160

    def sheet_inputs(self):
        ws = self.wb.create_sheet("Inputs")
        header(ws, 1, ["Name", "Value", "Unit", "Description", "Source / note"], [24, 16, 16, 70, 90])
        r = 2
        ws.cell(row=r, column=1, value="Scenario levers").font = BOLD
        r += 1
        lever_desc = {
            "s_bh": ("0/1", "1 = business hours only (~220 h/month); 0 = 24/7", 0),
            "s_itc": ("0/1", "1 = company claims GST input tax credit", 1),
            "s_spare": ("0/1", "1 = one spare replica per model (N+1)", 0),
            "s_years": ("years", "TCO horizon (3 or 5)", 3),
        }
        for nm, (unit, desc, default) in lever_desc.items():
            ws.cell(row=r, column=1, value=nm).font = BOLD
            c = ws.cell(row=r, column=2, value=default)
            c.font, c.fill = BLUE, YELLOW
            ws.cell(row=r, column=3, value=unit).font = BLACK
            ws.cell(row=r, column=4, value=desc).font = BLACK
            name(self.wb, nm, "Inputs", ("B", r))
            r += 1
        r += 1
        groups = {}
        for k in self.d.globals:
            groups.setdefault(k.split("_")[0], []).append(k)
        titles = {"s": "Scenario", "p": "Cloud extras", "t": "Tax & import", "h": "Hardware extras (USD)",
                  "f": "Facility: own room in Baner", "o": "Colocation", "k": "Finance", "e": "Serving defaults",
                  "q": "OCR"}
        for g, keys in groups.items():
            ws.cell(row=r, column=1, value=titles.get(g, g)).font = BOLD
            ws.cell(row=r, column=1).fill = GROUP_FILL
            r += 1
            for k in keys:
                m = meta(self.d.globals[k])
                unit, desc = (INPUTS.get(k) or OCR_INPUTS.get(k) or ("", ""))
                ws.cell(row=r, column=1, value=k).font = BLACK
                c = ws.cell(row=r, column=2, value=val(m))
                c.font, c.number_format = BLUE, numfmt(unit, k)
                ws.cell(row=r, column=3, value=unit).font = BLACK
                ws.cell(row=r, column=4, value=desc).font = BLACK
                ws.cell(row=r, column=5, value=src_text(m)).font = BLACK
                name(self.wb, k, "Inputs", ("B", r))
                self.sources.append(("Inputs", k, val(m), src_text(m)))
                r += 1
            r += 1
        ws.freeze_panes = "A2"

    def table_sheet(self, title, keyhdr, rows, cols, extra_cols=()):
        """rows: list of (key, {col: value-or-meta}); returns {col: column letter}, nrows."""
        ws = self.wb.create_sheet(title)
        allcols = list(extra_cols) + list(cols)
        header(ws, 1, [keyhdr] + allcols, [26] + [14] * len(allcols))
        units = ["unit"] + [(INPUTS.get(c) or OCR_INPUTS.get(c) or ("", ""))[0] for c in allcols]
        for i, u in enumerate(units, 1):
            ws.cell(row=2, column=i, value=u).font = Font(name=FONT, italic=True, color="595959")
        letters = {c: L(i + 2) for i, c in enumerate(allcols)}
        for r, (key, fields) in enumerate(rows, 3):
            ws.cell(row=r, column=1, value=key).font = BOLD
            for c in allcols:
                if c not in fields:
                    continue
                v = fields[c]
                cell = ws.cell(row=r, column=list(letters).index(c) + 2)
                if isinstance(v, str) and v.startswith("="):
                    cell.value, cell.font = v, BLACK
                else:
                    m = meta(v)
                    cell.value, cell.font = val(m), BLUE
                    s = src_text(m)
                    if s:
                        cell.comment = Comment(s, "cost model")
                        self.sources.append((title, f"{key} / {c}", val(m), s))
                unit = (INPUTS.get(c) or OCR_INPUTS.get(c) or ("", ""))[0]
                cell.number_format = numfmt(unit, c)
        ws.freeze_panes = "B3"
        return letters, len(rows)

    def sheet_tables(self):
        d = self.d
        # Models
        mcols = sorted({k for m in d.models.values() for p in m["precisions"].values() for k in p if k.startswith("m_")}
                       | {k for m in d.models.values() for k in m if k.startswith("m_")})
        rows = []
        for mk, m in d.models.items():
            for pk, p in m["precisions"].items():
                f = {k: v for k, v in m.items() if k.startswith("m_")}
                f.update({k: v for k, v in p.items() if k.startswith("m_")})
                f["compute_fmt"] = p["compute_fmt"]
                rows.append((f"{mk}|{pk}", f))
        self.mL, self.mN = self.table_sheet("Models", "model|precision", rows, mcols, ["compute_fmt"])
        # GPUs
        gcols = sorted({k for g in d.gpus.values() for k in g if k.startswith("g_")})
        fmts = ["fp4", "fp8", "bf16", "fp32"]
        rows = []
        for gk, g in d.gpus.items():
            f = {k: v for k, v in g.items() if k.startswith("g_")}
            for fm in fmts:
                tv = g["tflops"].get(fm)
                f[f"tflops_{fm}"] = tv if tv is not None else None
            rows.append((gk, f))
        self.gL, self.gN = self.table_sheet("GPUs", "gpu", rows, gcols + [f"tflops_{x}" for x in fmts])
        ws = self.wb["GPUs"]
        # FP4 fallback (Hopper): formula pointing at the fallback format, so edits flow through
        for r, (gk, g) in enumerate(rows, 3):
            if self.d.gpus[gk]["tflops"].get("fp4") is None:
                fb = self.d.gpus[gk].get("fp4_fallback", "fp8")
                c = ws[f"{self.gL['tflops_fp4']}{r}"]
                c.value, c.font = f"={self.gL['tflops_' + fb]}{r}", BLACK
                c.comment = Comment(f"No FP4 tensor cores: FP4 weights run as W4A8 on {fb.upper()} cores.", "cost model")
        # Efficiency (fully expanded)
        ecols = sorted(d.eff_default)
        rows = []
        for mk, m in d.models.items():
            for gk in d.gpus:
                for pk in m["precisions"]:
                    e = dict(d.eff_default)
                    for key in (mk, gk, f"{mk}/{gk}", f"{mk}/{gk}/{pk}"):
                        e.update(d.eff_override.get(key) or {})
                    rows.append((f"{mk}|{gk}|{pk}", e))
        self.eL, self.eN = self.table_sheet("Efficiency", "model|gpu|precision", rows, ecols)
        # Workloads
        wcols = sorted({k for w in d.workloads.values() for k in w if k[:2] in ("w_", "c_")})
        rows = [(wk, {k: v for k, v in w.items() if k[:2] in ("w_", "c_")}) for wk, w in d.workloads.items()]
        self.wL, self.wN = self.table_sheet("Workloads", "workload", rows, wcols)

    # -------------------------------------------------------------- lookups
    def lookup(self, var, keycells):
        """Excel formula fetching row-varying input `var` for the given key cells."""
        mk, pk, gk, wk, uk, sk = keycells
        rng = lambda sh, col, n: f"'{sh}'!${col}$3:${col}${n + 2}"
        if var == "s_users":
            return f"={uk}"
        if var == "s_speed":
            return f"={sk}"
        if var.startswith("m_"):
            return f"=INDEX({rng('Models', self.mL[var], self.mN)},MATCH({mk}&\"|\"&{pk},{rng('Models', 'A', self.mN)},0))"
        if var == "g_tflops":
            fmt = f"INDEX({rng('Models', self.mL['compute_fmt'], self.mN)},MATCH({mk}&\"|\"&{pk},{rng('Models', 'A', self.mN)},0))"
            first, last = self.gL["tflops_fp4"], self.gL["tflops_fp32"]
            return (f"=INDEX('GPUs'!${first}$3:${last}${self.gN + 2},MATCH({gk},{rng('GPUs', 'A', self.gN)},0),"
                    f"MATCH(\"tflops_\"&{fmt},'GPUs'!${first}$1:${last}$1,0))")
        if var.startswith("g_"):
            return f"=INDEX({rng('GPUs', self.gL[var], self.gN)},MATCH({gk},{rng('GPUs', 'A', self.gN)},0))"
        if var.startswith("e_"):
            return f"=INDEX({rng('Efficiency', self.eL[var], self.eN)},MATCH({mk}&\"|\"&{gk}&\"|\"&{pk},{rng('Efficiency', 'A', self.eN)},0))"
        if var[:2] in ("w_", "c_"):
            return f"=INDEX({rng('Workloads', self.wL[var], self.wN)},MATCH({wk},{rng('Workloads', 'A', self.wN)},0))"
        raise KeyError(var)

    # ------------------------------------------------------------ calculator
    def sheet_calculator(self):
        ws = self.wb.create_sheet("Calculator")
        header(ws, 1, ["Variable", "Value", "Unit", "Description", "Formula"], [22, 18, 16, 70, 80])
        d = self.d
        sel_vals = {"sel_model": next(iter(d.models)), "sel_precision": "native", "sel_gpu": next(iter(d.gpus)),
                    "sel_workload": next(iter(d.workloads)), "sel_users": 90, "sel_speed": 30}
        ref = d.scen.get("calculator_default") or {}
        for k in SEL:
            if k[4:] in ref:
                sel_vals[k] = ref[k[4:]]
        lists = {"sel_model": ",".join(d.models), "sel_precision": ",".join(
            sorted({p for m in d.models.values() for p in m["precisions"]})),
                 "sel_gpu": ",".join(d.gpus), "sel_workload": ",".join(d.workloads)}
        r = 2
        ws.cell(row=r, column=1, value="Selection").font = BOLD
        r += 1
        for k in SEL:
            ws.cell(row=r, column=1, value=k).font = BOLD
            c = ws.cell(row=r, column=2, value=sel_vals[k])
            c.font, c.fill = BLUE, YELLOW
            if k in lists:
                dv = DataValidation(type="list", formula1=f'"{lists[k]}"', allow_blank=False)
                ws.add_data_validation(dv)
                dv.add(c)
            name(self.wb, k, "Calculator", ("B", r))
            r += 1
        r += 1
        ws.cell(row=r, column=1, value="Looked-up inputs").font = BOLD
        r += 1
        keycells = SEL
        for var in self.row_inputs:
            unit, desc = INPUTS[var]
            ws.cell(row=r, column=1, value=var)
            c = ws.cell(row=r, column=2, value=self.lookup(var, keycells))
            c.font, c.number_format = GREEN, numfmt(unit, var)
            ws.cell(row=r, column=3, value=unit)
            ws.cell(row=r, column=4, value=desc)
            name(self.wb, var, "Calculator", ("B", r))
            r += 1
        group = None
        for f in self.llm_order:
            if f.group != group:
                r += 1
                ws.cell(row=r, column=1, value=f.group.title()).font = BOLD
                ws.cell(row=r, column=1).fill = GROUP_FILL
                group = f.group
                r += 1
            ws.cell(row=r, column=1, value=f.name)
            c = ws.cell(row=r, column=2, value=to_xl(f.expr, lambda n: n))
            c.font, c.number_format = BLACK, numfmt(f.unit, f.name)
            ws.cell(row=r, column=3, value=f.unit)
            ws.cell(row=r, column=4, value=f.desc)
            ws.cell(row=r, column=5, value=f.expr).font = Font(name=FONT, color="7F7F7F")
            name(self.wb, f.name, "Calculator", ("B", r))
            r += 1
        for row in ws.iter_rows(min_row=2, max_row=r, max_col=5):
            for c in row:
                if c.font.name != FONT:
                    c.font = Font(name=FONT, bold=c.font.bold, color=c.font.color)
        ws.freeze_panes = "A2"

    # ---------------------------------------------------------------- matrix
    def matrix_rows(self):
        if self.rows_mode == "reference":
            for sid, s in self.d.scen["reference"].items():
                yield {k: s[k] for k in ROW_KEYS}
        else:
            yield from self.d.llm_grid()

    def sheet_matrix(self):
        ws = self.wb.create_sheet("Matrix")
        keys = list(ROW_KEYS)
        cols = keys + self.row_inputs + [f.name for f in self.llm_order] + ["grp", "best_buy", "best_rent"]
        col = {c: L(i + 1) for i, c in enumerate(cols)}
        self.xL = col
        groups = ["key"] * len(keys) + ["input"] * len(self.row_inputs) + [f.group for f in self.llm_order] + ["summary"] * 3
        for i, (c, g) in enumerate(zip(cols, groups), 1):
            ws.cell(row=1, column=i, value=g).font = BOLD
            h = ws.cell(row=2, column=i, value=c)
            h.font, h.fill = HEAD, HEAD_FILL
            unit = (INPUTS.get(c) or ("",))[0] if c in INPUTS else next((f.unit for f in self.llm_order if f.name == c), "")
            ws.cell(row=3, column=i, value=unit).font = Font(name=FONT, italic=True, color="595959")
            ws.column_dimensions[L(i)].width = 12
        rows = list(self.matrix_rows())
        n0, n1 = 4, 4 + len(rows) - 1
        self.matrix_span = (n0, n1)
        for r, k in enumerate(rows, n0):
            for kk in keys:
                ws[f"{col[kk]}{r}"] = k[kk]
                ws[f"{col[kk]}{r}"].font = BLUE
            kc = tuple(f"${col[kk]}{r}" for kk in keys)
            for var in self.row_inputs:
                c = ws[f"{col[var]}{r}"]
                c.value, c.font = self.lookup(var, kc), GREEN
            def ref(n, r=r):
                return n if n in self.global_names else f"{col[n]}{r}"
            for f in self.llm_order:
                c = ws[f"{col[f.name]}{r}"]
                c.value, c.font = to_xl(f.expr, ref), BLACK
                c.number_format = numfmt(f.unit, f.name)
            # summary helpers: cheapest feasible option per (model, workload, users, speed)
            g = f"{col['model']}{r}&\"|\"&{col['workload']}{r}&\"|\"&{col['users']}{r}&\"|\"&{col['speed']}{r}"
            ws[f"{col['grp']}{r}"] = f"={g}"
            rngc = lambda c: f"${col[c]}${n0}:${col[c]}${n1}"
            for out, metric in (("best_buy", "buy_m"), ("best_rent", "rent_best")):
                mn = (f"_xlfn.MINIFS({rngc(metric)},{rngc('grp')},{col['grp']}{r},{rngc('feasible')},1)")
                ws[f"{col[out]}{r}"] = (f"=IF(AND({col['feasible']}{r}=1,{col[metric]}{r}={mn}),"
                                        f"{col['grp']}{r}&\"|1\",\"\")")
        ws.freeze_panes = f"{L(len(keys) + 1)}4"

    def sheet_summary(self):
        ws = self.wb.create_sheet("Summary")
        header(ws, 1, ["Model", "Workload", "Users", "Speed tok/s", "Cheapest own-room ₹/month", "GPU", "Precision",
                       "Nodes", "Cheapest rent ₹/month", "GPU (rent)", "Precision (rent)", "Best colo ₹/month",
                       "Own room per user ₹/month"], [16, 12, 8, 10, 20, 10, 10, 8, 20, 10, 12, 18, 18])
        n0, n1 = self.matrix_span
        col = self.xL
        rng = lambda c: f"Matrix!${col[c]}${n0}:${col[c]}${n1}"
        ax = self.d.scen["grid"]
        r = 2
        for m in ax["models"]:
            for w in ax["workloads"]:
                for u in ax["users"]:
                    for s in ax["speeds"]:
                        for i, v in enumerate((m, w, u, s), 1):
                            ws.cell(row=r, column=i, value=v).font = BLUE
                        g = f"A{r}&\"|\"&B{r}&\"|\"&C{r}&\"|\"&D{r}"
                        crit = f"{rng('grp')},{g},{rng('feasible')},1"
                        ws[f"E{r}"] = f"=IFERROR(_xlfn.MINIFS({rng('buy_m')},{crit}),\"\")"
                        mb = f"MATCH({g}&\"|1\",{rng('best_buy')},0)"
                        ws[f"F{r}"] = f"=IFERROR(INDEX({rng('gpu')},{mb}),\"none feasible\")"
                        ws[f"G{r}"] = f"=IFERROR(INDEX({rng('precision')},{mb}),\"\")"
                        ws[f"H{r}"] = f"=IFERROR(INDEX({rng('nodes')},{mb}),\"\")"
                        ws[f"I{r}"] = f"=IFERROR(_xlfn.MINIFS({rng('rent_best')},{crit}),\"\")"
                        mr = f"MATCH({g}&\"|1\",{rng('best_rent')},0)"
                        ws[f"J{r}"] = f"=IFERROR(INDEX({rng('gpu')},{mr}),\"\")"
                        ws[f"K{r}"] = f"=IFERROR(INDEX({rng('precision')},{mr}),\"\")"
                        ws[f"L{r}"] = f"=IFERROR(_xlfn.MINIFS({rng('colo_m')},{crit}),\"\")"
                        ws[f"M{r}"] = f"=IFERROR(E{r}/C{r},\"\")"
                        for c in "EILM":
                            ws[f"{c}{r}"].number_format = '"₹"#,##0'
                        r += 1
        for row in ws.iter_rows(min_row=2, max_row=r):
            for c in row:
                if c.font.color is None or c.font.color.rgb in (None, "FF000000"):
                    c.font = BLACK
        ws.freeze_panes = "E2"

    # ------------------------------------------------------------------- OCR
    def sheet_ocr(self):
        d = self.d
        o = d.ocr
        qcols = sorted({k for g in o["gpus"].values() for k in g if k.startswith("q_")}
                       | {k for g in o["gpus"].values() for p in g["precisions"].values() for k in p if k.startswith("q_")})
        rows = []
        for gk, g in o["gpus"].items():
            for pk, p in g["precisions"].items():
                f = {k: v for k, v in g.items() if k.startswith("q_")}
                f.update({k: v for k, v in p.items() if k.startswith("q_")})
                rows.append((f"{gk}|{pk}", f))
        qL, qN = self.table_sheet("OCR inputs", "gpu|precision", rows, qcols)
        ws = self.wb["OCR inputs"]
        tr = qN + 5
        ws.cell(row=tr, column=1, value="Load tiers (pages/hour at peak)").font = BOLD
        tiers = {}
        for i, (t, v) in enumerate(o["tiers"].items(), tr + 1):
            ws.cell(row=i, column=1, value=t).font = BOLD
            ws.cell(row=i, column=2, value=v).font = BLUE
            tiers[t] = f"'OCR inputs'!$B${i}"

        ws = self.wb.create_sheet("OCR")
        keys = ["gpu", "precision", "tier"]
        row_in = [k for k in OCR_INPUTS if k not in self.global_names]
        cols = keys + row_in + [f.name for f in self.ocr_order]
        col = {c: L(i + 1) for i, c in enumerate(cols)}
        for i, c in enumerate(cols, 1):
            h = ws.cell(row=1, column=i, value=c)
            h.font, h.fill = HEAD, HEAD_FILL
            ws.column_dimensions[L(i)].width = 12
        for r, k in enumerate(d.ocr_grid(), 2):
            for kk in keys:
                ws[f"{col[kk]}{r}"] = k[kk]
                ws[f"{col[kk]}{r}"].font = BLUE
            key = f"${col['gpu']}{r}&\"|\"&${col['precision']}{r}"
            for var in row_in:
                c = ws[f"{col[var]}{r}"]
                if var == "q_pages_h":
                    c.value = f"=INDEX('OCR inputs'!$B${tr + 1}:$B${tr + len(o['tiers'])},MATCH(${col['tier']}{r},'OCR inputs'!$A${tr + 1}:$A${tr + len(o['tiers'])},0))"
                else:
                    c.value = (f"=INDEX('OCR inputs'!${qL[var]}$3:${qL[var]}${qN + 2},"
                               f"MATCH({key},'OCR inputs'!$A$3:$A${qN + 2},0))")
                c.font = GREEN

            def ref(n, r=r):
                return n if n in self.global_names else f"{col[n]}{r}"
            for f in self.ocr_order:
                c = ws[f"{col[f.name]}{r}"]
                c.value, c.font = to_xl(f.expr, ref), BLACK
                c.number_format = numfmt(f.unit, f.name)
        ws.freeze_panes = "D2"

    def sheet_sources(self):
        ws = self.wb.create_sheet("Sources")
        header(ws, 1, ["Sheet", "Item", "Value", "Source | date | note"], [14, 40, 16, 140])
        for r, (sh, item, v, s) in enumerate(self.sources, 2):
            for i, x in enumerate((sh, item, v, s or "ASSUMPTION (no published source)"), 1):
                ws.cell(row=r, column=i, value=x).font = BLACK

    def build(self, path):
        self.sheet_readme()
        self.sheet_inputs()
        self.sheet_tables()
        self.sheet_calculator()
        self.sheet_matrix()
        if self.rows_mode == "full":
            self.sheet_summary()
        self.sheet_ocr()
        self.sheet_sources()
        order = ["Read me", "Summary", "Calculator", "Inputs", "Matrix", "OCR", "Models", "GPUs", "Efficiency",
                 "Workloads", "OCR inputs", "Sources"]
        self.wb._sheets = [self.wb[s] for s in order if s in self.wb.sheetnames]
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.save(path)
        return path


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 and not sys.argv[1].startswith("--") else \
        Path(__file__).parent.parent / "reports" / "Self hosting LLM cost Pune.xlsx"
    mode = "reference" if "--rows" in sys.argv and sys.argv[sys.argv.index("--rows") + 1] == "reference" else "full"
    Builder(Data(), mode).build(out)
    print(f"wrote {out} ({mode})")


if __name__ == "__main__":
    main()
